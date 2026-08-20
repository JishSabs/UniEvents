#!/usr/bin/env python3
"""
Generate project documentation and presentation for UniEvents.
Creates outputs/UniEvents_Documentation.docx, UniEvents_Presentation.pptx,
outputs/diagrams/*.png|.svg and outputs/package.zip

Usage:
  python scripts/generate_docs.py

Requires: see scripts/requirements.txt
"""
import os
import re
import sys
from pathlib import Path
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs"
DIAG = OUT / "diagrams"
ASSETS = OUT / "assets"
OUT.mkdir(exist_ok=True)
DIAG.mkdir(exist_ok=True)
ASSETS.mkdir(parents=True, exist_ok=True)

# Files to include in appendix
FILES_TO_EMBED = [
    ROOT / "PROJECT_DOCUMENTATION.md",
    ROOT / "firestore.rules",
    ROOT / "firestore.indexes.json",
    ROOT / "src/types/index.ts",
    ROOT / "src/lib/transactions.ts",
]

# Helper: simple TypeScript interface parser (very tolerant)
INTERFACE_RE = re.compile(r"export interface\s+(\w+)\s*{([\s\S]*?)}", re.M)
FIELD_RE = re.compile(r"^\s*([\w?]+):\s*([^;]+);", re.M)


def parse_types(ts_path: Path):
    if not ts_path.exists():
        return {}
    txt = ts_path.read_text(encoding="utf-8")
    out = {}
    for m in INTERFACE_RE.finditer(txt):
        name = m.group(1)
        body = m.group(2)
        fields = []
        for f in FIELD_RE.finditer(body):
            fname = f.group(1).strip()
            ftype = f.group(2).strip()
            fields.append((fname, ftype))
        out[name] = fields
    return out


def write_dot(er_dot: str, dot_path: Path):
    dot_path.write_text(er_dot, encoding="utf-8")
    # Try rendering with graphviz if available
    try:
        from graphviz import Source

        s = Source(er_dot)
        png = str(dot_path.with_suffix(".png"))
        svg = str(dot_path.with_suffix(".svg"))
        s.format = "png"
        s.render(filename=str(dot_path.with_suffix("")), cleanup=True)
        # graphviz may also produce svg; attempt explicitly
        s.format = "svg"
        s.render(filename=str(dot_path.with_suffix("")), cleanup=True)
        print(f"Rendered {png} and {svg}")
        return dot_path.with_suffix(".png"), dot_path.with_suffix(".svg")
    except Exception as e:
        print("Graphviz render failed or not available:", e)
        return None, None


def build_er_diagram(parsed_types: dict) -> str:
    # Build a conservative ER diagram focusing on collections in this app
    nodes = []
    edges = []

    def node_label(name, fields):
        # show a few key fields
        sample = "\\l".join([f"{n}: {t}" for n, t in fields[:6]]) + "\\l"
        return f"{name}|{sample}"

    mapping = {
        "UserProfile": "Users",
        "Announcement": "Announcements",
        "MarketplaceListing": "Listings",
        "Transaction": "Transactions",
        "Message": "Messages",
    }

    for itype, label in mapping.items():
        fields = parsed_types.get(itype, [])
        if not fields:
            # fallback to a small shape
            nodes.append(f'  {label} [label="{{{label}|id: string\\lid: ...\\l}}", shape=record];')
        else:
            lbl = node_label(label, fields)
            nodes.append(f'  {label} [label="{{{lbl}}}", shape=record];')

    # Edges for common relations
    edges.append('  Users -> Announcements [label="author"];')
    edges.append('  Users -> Listings [label="author"];')
    edges.append('  Users -> Transactions [label="buyer/seller"];')
    edges.append('  Transactions -> Messages [label="has subcollection"];')

    dot = ["digraph UniEvents {", '  node [shape=record,fontname="Helvetica"];'] + nodes + edges + ["}"]
    return "\n".join(dot)


def create_docx(docx_path: Path, parsed_types: dict, files: list, er_png: Path = None, arch_png: Path = None):
    try:
        from docx import Document
        from docx.shared import Inches
    except Exception as e:
        print("python-docx is not installed. Install with: pip install python-docx")
        raise

    doc = Document()
    doc.core_properties.title = "UniEvents Documentation"
    doc.add_heading("UniEvents - System Documentation", level=1)
    doc.add_paragraph("Generated documentation. Update placeholders as needed.")

    doc.add_heading("Executive Summary", level=2)
    doc.add_paragraph("UniEvents is a small campus marketplace and announcements platform built with Next.js and Firebase.")

    doc.add_heading("Architecture", level=2)
    if arch_png and arch_png.exists():
        doc.add_picture(str(arch_png), width=Inches(6))
    else:
        doc.add_paragraph("[Architecture diagram placeholder]")

    doc.add_heading("Data Model (ER)", level=2)
    if er_png and er_png.exists():
        doc.add_picture(str(er_png), width=Inches(6))
    else:
        doc.add_paragraph("[ER diagram placeholder]")

    doc.add_heading("Collections and Types", level=2)
    for name, fields in parsed_types.items():
        doc.add_heading(name, level=3)
        tbl = doc.add_table(rows=1, cols=2)
        hdr = tbl.rows[0].cells
        hdr[0].text = "Field"
        hdr[1].text = "Type"
        for f, t in fields:
            r = tbl.add_row().cells
            r[0].text = f
            r[1].text = t

    doc.add_heading("Appendix: Important files", level=2)
    for p in files:
        if p.exists():
            doc.add_heading(p.name, level=3)
            content = p.read_text(encoding="utf-8")
            # limit size
            excerpt = content[:4000]
            doc.add_paragraph(excerpt)
            if len(content) > len(excerpt):
                doc.add_paragraph("... (truncated) ...")

    doc.save(str(docx_path))
    print(f"Wrote {docx_path}")


def create_pptx(pptx_path: Path, er_png: Path = None, arch_png: Path = None):
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except Exception as e:
        print("python-pptx is not installed. Install with: pip install python-pptx")
        raise

    prs = Presentation()
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "UniEvents"
    subtitle.text = "System overview and diagrams"

    # Architecture slide
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    title = slide.shapes.title
    title.text = "Architecture"
    if arch_png and arch_png.exists():
        slide.shapes.add_picture(str(arch_png), Inches(1), Inches(1.5), width=Inches(8))
    else:
        tx = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(8), Inches(2))
        tx.text_frame.text = "[Architecture diagram placeholder]"

    # ER slide
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    title = slide.shapes.title
    title.text = "Data Model (ER)"
    if er_png and er_png.exists():
        slide.shapes.add_picture(str(er_png), Inches(1), Inches(1.5), width=Inches(8))
    else:
        tx = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(8), Inches(2))
        tx.text_frame.text = "[ER diagram placeholder]"

    prs.save(str(pptx_path))
    print(f"Wrote {pptx_path}")


def package_outputs(out_dir: Path, zip_path: Path):
    with ZipFile(zip_path, "w") as z:
        for p in out_dir.rglob("*"):
            if p.is_file():
                z.write(p, p.relative_to(out_dir))
    print(f"Packaged outputs to {zip_path}")


def main():
    print("Parsing types...")
    parsed = parse_types(ROOT / "src/types/index.ts")

    er_dot = build_er_diagram(parsed)
    dot_path = DIAG / "er_diagram.dot"
    dot_path.write_text(er_dot, encoding="utf-8")
    print(f"Wrote DOT to {dot_path}")

    er_png, er_svg = write_dot(er_dot, DIAG / "er_diagram")

    # architecture diagram (simple)
    arch_dot = 'digraph arch { node [shape=box]; WebApp -> FirebaseAuth; WebApp -> Firestore; WebApp -> RealtimeDB; WebApp -> Storage; }'
    (DIAG / "architecture.dot").write_text(arch_dot)
    arch_png, arch_svg = write_dot(arch_dot, DIAG / "architecture")

    # create docx and pptx
    docx_path = OUT / "UniEvents_Documentation.docx"
    pptx_path = OUT / "UniEvents_Presentation.pptx"

    try:
        create_docx(docx_path, parsed, [p for p in FILES_TO_EMBED if p.exists()], er_png, arch_png)
    except Exception as e:
        print("Skipping DOCX generation:", e)

    try:
        create_pptx(pptx_path, er_png, arch_png)
    except Exception as e:
        print("Skipping PPTX generation:", e)

    # copy any screenshots if present
    screenshots = list((ROOT / "public").glob("**/*.png")) if (ROOT / "public").exists() else []
    for s in screenshots:
        dest = ASSETS / s.name
        try:
            dest.write_bytes(s.read_bytes())
        except Exception:
            pass

    package_outputs(OUT, OUT / "package.zip")


if __name__ == "__main__":
    main()
